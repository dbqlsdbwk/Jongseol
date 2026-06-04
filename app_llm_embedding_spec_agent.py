# -*- coding: utf-8 -*-
"""
건설 시방서 검토 통합 UI
PDF → TXT 추출 → Rule Agent → Embedding Retrieval Agent → LLM Agent → 보고서 생성

실행:
    pip install streamlit pandas openpyxl pdfplumber pymupdf python-docx openai numpy
    streamlit run app_llm_embedding_spec_agent.py

API KEY 설정 방법:
    1) 로컬 환경변수
       setx OPENAI_API_KEY "sk-..."
    2) Streamlit Cloud secrets
       OPENAI_API_KEY = "sk-..."
    3) 앱 사이드바에 직접 입력
"""

import os
import re
import io
import json
import time
import tempfile
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None

try:
    from openai import OpenAI
except Exception:
    OpenAI = None

try:
    from docx import Document
    from docx.shared import Pt
except Exception:
    Document = None


# =========================================================
# 1. 기본 Rule 정의
# =========================================================

DEFAULT_TRADE_RULES = [
    {
        "trade": "방수공사",
        "target_keywords": ["방수", "도막", "우레탄", "시트방수"],
        "required": {
            "재료명": r"(우레탄|시트|도막|방수재|프라이머)",
            "두께 기준": r"\d+(\.\d+)?\s*(mm|㎜)",
            "시공 횟수": r"\d+\s*(회|차)",
            "양생/건조 시간": r"\d+\s*(시간|일)",
            "검사/시험 기준": r"(검사|시험|KS|품질)"
        }
    },
    {
        "trade": "지붕공사",
        "target_keywords": ["지붕", "마감재", "패널", "금속판"],
        "required": {
            "마감재 종류": r"(마감재|패널|금속판|징크|칼라강판)",
            "고정 방법": r"(고정|볼트|피스|체결)",
            "겹침/이음 기준": r"(겹침|이음|접합)",
            "배수/구배 기준": r"(배수|구배|물매)",
            "검사 기준": r"(검사|확인|품질)"
        }
    },
    {
        "trade": "외벽보수",
        "target_keywords": ["외벽", "보수", "균열", "박리", "누수"],
        "required": {
            "보수 범위": r"(범위|구간|위치|면적)",
            "손상 조사": r"(균열|박리|누수|손상)",
            "보수 재료": r"(보수재|실링재|몰탈|모르타르|에폭시)",
            "시공 방법": r"(주입|충전|도포|제거|보수)",
            "완료 검사": r"(검사|확인|점검)"
        }
    },
    {
        "trade": "천정공사",
        "target_keywords": ["천정", "천장", "보강", "텍스", "석고보드"],
        "required": {
            "마감재 종류": r"(마감재|텍스|석고보드|패널)",
            "보강 기준": r"(보강|보강재|철물)",
            "고정 방법": r"(고정|체결|피스|앵커)",
            "처짐/변형 확인": r"(처짐|변형|수평)",
            "검사 기준": r"(검사|확인|점검)"
        }
    },
    {
        "trade": "폐기물처리",
        "target_keywords": ["폐기물", "철거물", "반출", "폐자재"],
        "required": {
            "분리/수거 기준": r"(분리|수거|선별)",
            "반출 기준": r"(반출|운반)",
            "처리 방법": r"(처리|폐기|재활용)",
            "관련 법령": r"(법령|폐기물관리법|관련법)",
            "기록/확인": r"(기록|확인|계근|인계서)"
        }
    },
    {
        "trade": "도장공사",
        "target_keywords": ["도장", "페인트", "도료", "칠"],
        "required": {
            "도료 종류": r"(도료|페인트|프라이머|상도|하도)",
            "바탕 처리": r"(바탕|면처리|연마|청소)",
            "도장 횟수": r"\d+\s*(회|차)",
            "건조 시간": r"\d+\s*(시간|일)",
            "검사 기준": r"(검사|확인|품질)"
        }
    },
    {
        "trade": "철거공사",
        "target_keywords": ["철거", "해체", "제거"],
        "required": {
            "철거 범위": r"(범위|구간|위치|대상)",
            "철거 방법": r"(철거|해체|절단|제거)",
            "안전 조치": r"(안전|보호|가설|방호)",
            "폐기물 처리": r"(폐기물|반출|처리)",
            "확인/검사": r"(확인|검사|점검)"
        }
    },
    {
        "trade": "콘크리트공사",
        "target_keywords": ["콘크리트", "타설", "압축강도"],
        "required": {
            "강도 기준": r"(\d+\s*(MPa|N/mm2|㎫)|압축강도)",
            "타설 방법": r"(타설|다짐|진동)",
            "양생 기준": r"(양생|\d+\s*(일|시간))",
            "시험 기준": r"(시험|공시체|KS)",
            "허용오차": r"(허용오차|오차)"
        }
    },
    {
        "trade": "철근공사",
        "target_keywords": ["철근", "배근", "이음", "정착"],
        "required": {
            "철근 종류/규격": r"(철근|D\d+|SD\d+)",
            "배근 기준": r"(배근|간격|피복)",
            "이음 기준": r"(이음|겹침|정착)",
            "검사 기준": r"(검사|확인|검측)",
            "도면 기준": r"(도면|상세도)"
        }
    },
    {
        "trade": "안전관리",
        "target_keywords": ["안전", "보호구", "추락", "위험"],
        "required": {
            "작업자 보호": r"(보호구|안전모|안전화|안전대)",
            "추락 방지": r"(추락|난간|안전망|개구부)",
            "위험성 관리": r"(위험|위험성|작업계획)",
            "교육/점검": r"(교육|점검|확인)",
            "책임 주체": r"(시공자|관리자|감독자|책임)"
        }
    }
]

DEFAULT_VAGUE_WORDS = [
    "적절히", "충분히", "필요시", "관련 기준에 따라",
    "감독자 지시에 따라", "양호하게", "견고히",
    "깨끗이", "이상 없도록", "적정하게",
    "가능한", "필요한", "상당한", "적합하게"
]

GENERAL_EXCLUDE_PATTERNS = [
    r"본 공사는",
    r"공사의 위치",
    r"적용범위",
    r"용어의 정의",
    r"발주자",
    r"수급인",
    r"설계서",
    r"제\d+장",
]


# =========================================================
# 2. 공통 유틸
# =========================================================

def now_name(prefix, ext):
    return f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"


def read_uploaded_text(uploaded_file):
    raw = uploaded_file.read()
    for enc in ["utf-8", "cp949", "euc-kr"]:
        try:
            return raw.decode(enc)
        except Exception:
            pass
    return raw.decode("utf-8", errors="ignore")


def clean_basic(text):
    text = re.sub(r"^===== Page \d+ =====$", "", str(text), flags=re.MULTILINE)
    text = re.sub(r"^Page\s*\d+\s*$", "", text, flags=re.MULTILINE | re.IGNORECASE)
    text = re.sub(r"^=+\s*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()


def merge_broken_lines(text):
    text = clean_basic(text)
    lines = text.splitlines()
    merged = []
    buffer = ""

    for line in lines:
        line = line.strip()
        if not line:
            continue

        if re.match(r"^제\s*\d+\s*장", line):
            if buffer:
                merged.append(buffer.strip())
                buffer = ""
            merged.append(line)
            continue

        if re.match(r"^\d+(\.\d+)*\s+[가-힣A-Za-z0-9·\-\(\) ]{1,70}$", line):
            if buffer:
                merged.append(buffer.strip())
                buffer = ""
            merged.append(line)
            continue

        if buffer:
            buffer += " " + line
        else:
            buffer = line

        if re.search(r"(다\.?|함\.?|음\.?|됨\.?|임\.?|한다\.?|하여야 한다\.?|하여야 함\.?)$", line):
            merged.append(buffer.strip())
            buffer = ""

    if buffer:
        merged.append(buffer.strip())

    text = "\n".join(merged)
    text = re.sub(r"\s+(\d+(?:\.\d+)+\s+)", r"\n\1", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_sentences_for_rule(text):
    text = clean_basic(text)
    pieces = re.split(r"(?<=[다함음됨임\.])\s+|\n+", text)

    rows = []
    current_section = "-"

    for p in pieces:
        p = p.strip()
        if not p:
            continue

        sec = re.match(r"^(\d+(?:\.\d+)+)", p)
        if sec:
            current_section = sec.group(1)

        if is_title_or_noise(p):
            continue

        rows.append({
            "조항번호": current_section,
            "문장": p
        })

    return rows


def split_sentences_for_reference(text):
    text = clean_basic(text)
    pieces = re.split(r"(?<=[다함음됨임\.])\s+|\n+", text)
    output = []
    current_section = "-"

    for p in pieces:
        p = p.strip()
        if not p:
            continue

        sec = re.match(r"^(\d+(?:\.\d+)+)", p)
        if sec:
            current_section = sec.group(1)

        if len(p) < 15:
            continue
        if re.match(r"^제\s*\d+\s*장", p):
            continue

        output.append({
            "기준조항": current_section,
            "기준문장": p
        })

    return output


def is_title_or_noise(s):
    s = str(s).strip()
    if len(s) < 10:
        return True
    if re.match(r"^\d+(\.\d+)*\s*[가-힣A-Za-z ]{1,40}$", s):
        return True
    if re.match(r"^제\s*\d+\s*장", s):
        return True
    return False


def make_download_bytes_df(df, file_type="xlsx"):
    buffer = io.BytesIO()
    if file_type == "xlsx":
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="검토결과")
    else:
        df.to_csv(buffer, index=False, encoding="utf-8-sig")
    buffer.seek(0)
    return buffer


# =========================================================
# 3. PDF → TXT
# =========================================================

def extract_with_pdfplumber(file_bytes):
    if pdfplumber is None:
        raise RuntimeError("pdfplumber가 설치되어 있지 않습니다.")

    texts = []
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    try:
        with pdfplumber.open(tmp_path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                text = page.extract_text(x_tolerance=1, y_tolerance=3, layout=True)
                if text:
                    texts.append(f"\n===== Page {i} =====\n{text}")
    finally:
        try:
            os.remove(tmp_path)
        except Exception:
            pass

    return "\n".join(texts)


def extract_with_pymupdf(file_bytes):
    if fitz is None:
        raise RuntimeError("pymupdf가 설치되어 있지 않습니다.")

    texts = []
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text")
        if text:
            texts.append(f"\n===== Page {i} =====\n{text}")
    return "\n".join(texts)


def pdf_to_clean_txt(uploaded_pdf, method):
    file_bytes = uploaded_pdf.read()
    if method == "pdfplumber":
        raw_text = extract_with_pdfplumber(file_bytes)
    else:
        raw_text = extract_with_pymupdf(file_bytes)
    return merge_broken_lines(raw_text)


# =========================================================
# 4. Rule Agent
# =========================================================

def is_general_info(sentence):
    return any(re.search(pattern, str(sentence)) for pattern in GENERAL_EXCLUDE_PATTERNS)


def classify_trade(sentence, rules):
    matched = []
    for rule in rules:
        hit_keywords = [kw for kw in rule["target_keywords"] if kw and kw in sentence]
        if hit_keywords:
            matched.append((rule, hit_keywords))
    if not matched:
        return None, []
    matched.sort(key=lambda x: len(x[1]), reverse=True)
    return matched[0]


def inspect_required_items(sentence, rule):
    found = []
    missing = []
    for item, pattern in rule["required"].items():
        try:
            if re.search(pattern, sentence):
                found.append(item)
            else:
                missing.append(item)
        except re.error:
            missing.append(item)
    return found, missing


def detect_vague(sentence, vague_words):
    return [w for w in vague_words if w and w in sentence]


def judge(found, missing, vague):
    if not missing and not vague:
        return "양호"
    if vague and missing:
        return "누락 및 모호"
    if vague:
        return "모호 표현"
    if missing:
        return "누락 가능성"
    return "검토 필요"


def make_advice(trade, missing, vague):
    comments = []
    if missing:
        comments.append(f"{trade} 관련 문장이지만 {', '.join(missing)} 항목이 명확하지 않습니다.")
    if vague:
        comments.append(f"'{', '.join(vague)}' 표현은 구체적인 수치, 기준, 방법, 책임주체 또는 절차로 바꾸는 것이 좋습니다.")
    if not comments:
        comments.append("지정한 검토 기준에 비추어 필수 정보가 비교적 잘 포함되어 있습니다.")
    return " ".join(comments)


def analyze_spec_text(text, rules, vague_words):
    sentences = split_sentences_for_rule(text)
    results = []

    for row in sentences:
        section = row["조항번호"]
        sentence = row["문장"]
        vague = detect_vague(sentence, vague_words)

        if is_general_info(sentence):
            results.append({
                "조항번호": section,
                "문장": sentence,
                "분류": "일반정보/검토제외",
                "판정": "검토 제외",
                "누락 항목": "-",
                "모호 표현": ", ".join(vague) if vague else "-",
                "조언": "공종별 시공기준·품질기준·검사기준과 직접 관련성이 낮아 검토 대상에서 제외합니다.",
                "매칭 키워드": "-",
                "확인 항목": "-"
            })
            continue

        classified = classify_trade(sentence, rules)

        if classified[0] is None:
            if vague:
                results.append({
                    "조항번호": section,
                    "문장": sentence,
                    "분류": "기타/모호표현",
                    "판정": "모호 표현",
                    "누락 항목": "-",
                    "모호 표현": ", ".join(vague),
                    "조언": f"'{', '.join(vague)}' 표현은 구체적인 수치, 기준, 방법, 책임주체 또는 절차로 바꾸는 것이 좋습니다.",
                    "매칭 키워드": "-",
                    "확인 항목": "-"
                })
            else:
                results.append({
                    "조항번호": section,
                    "문장": sentence,
                    "분류": "기타 일반조항",
                    "판정": "검토 제외",
                    "누락 항목": "-",
                    "모호 표현": "-",
                    "조언": "현재 설정된 공종 기준과 직접 관련성이 낮습니다.",
                    "매칭 키워드": "-",
                    "확인 항목": "-"
                })
            continue

        rule, hit_keywords = classified
        found, missing = inspect_required_items(sentence, rule)
        status = judge(found, missing, vague)

        results.append({
            "조항번호": section,
            "문장": sentence,
            "분류": rule["trade"],
            "판정": status,
            "누락 항목": ", ".join(missing) if missing else "-",
            "모호 표현": ", ".join(vague) if vague else "-",
            "조언": make_advice(rule["trade"], missing, vague),
            "매칭 키워드": ", ".join(hit_keywords),
            "확인 항목": ", ".join(found) if found else "-"
        })

    return pd.DataFrame(results)


# =========================================================
# 5. OpenAI API Client
# =========================================================

def get_openai_api_key(sidebar_key=None):
    if sidebar_key:
        return sidebar_key.strip()

    env_key = os.getenv("OPENAI_API_KEY")
    if env_key:
        return env_key.strip()

    try:
        if "OPENAI_API_KEY" in st.secrets:
            return st.secrets["OPENAI_API_KEY"].strip()
    except Exception:
        pass

    return ""


def get_openai_client(api_key):
    if OpenAI is None:
        raise RuntimeError("openai 패키지가 설치되어 있지 않습니다. pip install openai 실행이 필요합니다.")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY가 설정되어 있지 않습니다.")
    return OpenAI(api_key=api_key)


# =========================================================
# 6. Embedding Retrieval Agent
# =========================================================

def load_reference_texts(reference_files):
    rows = []

    for f in reference_files:
        file_name = f.name
        text = read_uploaded_text(f)
        text = merge_broken_lines(text)
        sent_rows = split_sentences_for_reference(text)

        for r in sent_rows:
            rows.append({
                "근거파일": file_name,
                "기준조항": r["기준조항"],
                "기준문장": r["기준문장"]
            })

    return pd.DataFrame(rows)


def cosine_sim_matrix(query_vec, doc_matrix):
    q = np.array(query_vec, dtype=np.float32)
    docs = np.array(doc_matrix, dtype=np.float32)

    q_norm = np.linalg.norm(q)
    d_norm = np.linalg.norm(docs, axis=1)

    if q_norm == 0:
        return np.zeros(len(docs))

    denom = d_norm * q_norm
    denom[denom == 0] = 1e-12
    return np.dot(docs, q) / denom


def create_embeddings(client, texts, embedding_model="text-embedding-3-small", batch_size=64):
    vectors = []
    texts = [str(t)[:8000] for t in texts]

    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        resp = client.embeddings.create(
            model=embedding_model,
            input=batch
        )
        batch_vecs = [item.embedding for item in resp.data]
        vectors.extend(batch_vecs)

    return vectors


def build_reference_embedding_index(client, reference_df, embedding_model):
    if reference_df.empty:
        raise RuntimeError("기준 시방서 문장이 없습니다. 기준 TXT 파일을 업로드해주세요.")

    vectors = create_embeddings(
        client=client,
        texts=reference_df["기준문장"].tolist(),
        embedding_model=embedding_model
    )

    ref = reference_df.copy()
    ref["embedding"] = vectors
    return ref


def retrieve_evidence_for_sentence(client, sentence, ref_index, embedding_model, top_k=3, min_score=0.15):
    query_vec = create_embeddings(
        client=client,
        texts=[sentence],
        embedding_model=embedding_model
    )[0]

    doc_matrix = ref_index["embedding"].tolist()
    scores = cosine_sim_matrix(query_vec, doc_matrix)

    order = np.argsort(scores)[::-1]
    selected = []

    for idx in order[:top_k]:
        score = float(scores[idx])
        if score < min_score:
            continue

        row = ref_index.iloc[idx]
        selected.append({
            "근거파일": row["근거파일"],
            "기준조항": row["기준조항"],
            "기준문장": row["기준문장"],
            "유사도": round(score, 4)
        })

    return selected


def format_retrieval_evidence(evidences):
    if not evidences:
        return "-"

    lines = []
    for ev in evidences:
        lines.append(
            f"[{ev['근거파일']} / {ev['기준조항']} / 유사도 {ev['유사도']}] {ev['기준문장']}"
        )
    return "\n".join(lines)


def add_embedding_retrieval_to_df(
    client,
    rule_df,
    ref_index,
    embedding_model="text-embedding-3-small",
    top_k=3,
    max_rows=30
):
    target_status = ["누락 가능성", "모호 표현", "누락 및 모호", "검토 필요"]

    df = rule_df.copy()
    if "Retrieval 근거" not in df.columns:
        df["Retrieval 근거"] = "-"

    target_indices = df[df["판정"].isin(target_status)].head(max_rows).index.tolist()

    progress = st.progress(0)
    total = len(target_indices)

    for n, idx in enumerate(target_indices, start=1):
        sentence = str(df.at[idx, "문장"])
        evidences = retrieve_evidence_for_sentence(
            client=client,
            sentence=sentence,
            ref_index=ref_index,
            embedding_model=embedding_model,
            top_k=top_k
        )
        df.at[idx, "Retrieval 근거"] = format_retrieval_evidence(evidences)
        progress.progress(n / total if total else 1)

    progress.empty()
    return df


# =========================================================
# 7. LLM Agent
# =========================================================

def generate_llm_review(
    client,
    sentence,
    trade,
    status,
    missing_items,
    vague_words,
    advice,
    retrieval_evidence,
    generation_model="gpt-5.5"
):
    prompt = f"""
너는 건설 시방서 검토 보고서를 작성하는 Agent다.
아래 입력을 바탕으로 검토 의견, 보완 필요 사항, 수정 예시 문장을 작성하라.

[검토 대상 문장]
{sentence}

[공종 분류]
{trade}

[Rule 기반 판정]
{status}

[누락 항목]
{missing_items}

[모호 표현]
{vague_words}

[Rule 기반 조언]
{advice}

[Embedding Retrieval Agent가 찾은 유사 기준 문장]
{retrieval_evidence}

작성 형식:
1. 검토 의견:
2. 보완 필요 사항:
3. 수정 예시 문장:

작성 조건:
- 한국어로 작성한다.
- 보고서에 바로 넣을 수 있게 간결하고 공식적인 문체로 작성한다.
- Retrieval 근거에 없는 구체 수치, 법령명, 시험명은 임의로 만들지 않는다.
- 확실하지 않은 기준은 '해당 기준 확인 필요'라고 표현한다.
- 수정 예시 문장은 1개만 작성한다.
"""

    resp = client.responses.create(
        model=generation_model,
        input=prompt
    )
    return resp.output_text


def add_llm_reviews_to_df(
    client,
    df,
    generation_model="gpt-5.5",
    max_rows=10,
    sleep_sec=0.1
):
    target_status = ["누락 가능성", "모호 표현", "누락 및 모호", "검토 필요"]

    out = df.copy()
    if "LLM 검토 의견" not in out.columns:
        out["LLM 검토 의견"] = "-"

    target_indices = out[out["판정"].isin(target_status)].head(max_rows).index.tolist()

    progress = st.progress(0)
    total = len(target_indices)

    for n, idx in enumerate(target_indices, start=1):
        row = out.loc[idx]
        try:
            review = generate_llm_review(
                client=client,
                sentence=str(row.get("문장", "")),
                trade=str(row.get("분류", "")),
                status=str(row.get("판정", "")),
                missing_items=str(row.get("누락 항목", "-")),
                vague_words=str(row.get("모호 표현", "-")),
                advice=str(row.get("조언", "-")),
                retrieval_evidence=str(row.get("Retrieval 근거", "-")),
                generation_model=generation_model
            )
        except Exception as e:
            review = f"LLM 생성 실패: {e}"

        out.at[idx, "LLM 검토 의견"] = review
        progress.progress(n / total if total else 1)
        time.sleep(sleep_sec)

    progress.empty()
    return out


# =========================================================
# 8. 보고서 생성
# =========================================================

def make_summary_tables(df):
    status_count = df["판정"].value_counts().reset_index()
    status_count.columns = ["판정", "건수"]

    trade_count = df["분류"].value_counts().reset_index()
    trade_count.columns = ["분류", "건수"]

    problem_df = df[df["판정"].isin(["누락 가능성", "모호 표현", "누락 및 모호", "검토 필요"])].copy()

    return status_count, trade_count, problem_df


def generate_markdown_report(df, title="시방서 Rule·Retrieval·LLM 검토 보고서"):
    status_count, trade_count, problem_df = make_summary_tables(df)

    lines = []
    lines.append(f"# {title}")
    lines.append("")
    lines.append(f"- 생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    lines.append(f"- 전체 문장 수: {len(df)}개")
    lines.append(f"- 검토 필요 문장 수: {len(problem_df)}개")
    lines.append("")

    lines.append("## 1. 검토 개요")
    lines.append("본 보고서는 PDF 또는 TXT 형태의 시방서 문장을 대상으로 Rule Agent, Embedding Retrieval Agent, LLM Agent를 순차적으로 적용하여 생성하였다.")
    lines.append("Rule Agent는 공종별 키워드와 필수항목 포함 여부를 기준으로 누락 및 모호 표현을 1차 판정하였다.")
    lines.append("Embedding Retrieval Agent는 기준 시방서 문장을 벡터화하여 검토 대상 문장과 의미적으로 유사한 기준 문장을 검색하였다.")
    lines.append("LLM Agent는 Rule 판정과 Retrieval 근거를 바탕으로 검토 의견과 보완 예시 문장을 작성하였다.")
    lines.append("")

    lines.append("## 2. 판정 결과 요약")
    lines.append(status_count.to_markdown(index=False))
    lines.append("")

    lines.append("## 3. 공종별 분류 결과")
    lines.append(trade_count.head(15).to_markdown(index=False))
    lines.append("")

    lines.append("## 4. 대표 검토 필요 문장")
    show_cols = ["조항번호", "분류", "판정", "문장", "누락 항목", "모호 표현", "조언", "Retrieval 근거", "LLM 검토 의견"]
    show_cols = [c for c in show_cols if c in problem_df.columns]

    if problem_df.empty:
        lines.append("검토 필요 문장이 발견되지 않았다.")
    else:
        for _, row in problem_df.head(10).iterrows():
            lines.append(f"### 조항 {row.get('조항번호', '-')}")
            lines.append(f"- 분류: {row.get('분류', '-')}")
            lines.append(f"- 판정: {row.get('판정', '-')}")
            lines.append(f"- 대상 문장: {row.get('문장', '-')}")
            lines.append(f"- 누락 항목: {row.get('누락 항목', '-')}")
            lines.append(f"- 모호 표현: {row.get('모호 표현', '-')}")
            lines.append(f"- Rule 조언: {row.get('조언', '-')}")
            if "Retrieval 근거" in row:
                lines.append(f"- Retrieval 근거: {row.get('Retrieval 근거', '-')}")
            if "LLM 검토 의견" in row:
                lines.append("")
                lines.append(str(row.get("LLM 검토 의견", "-")))
            lines.append("")

    lines.append("## 5. 종합 의견")
    lines.append("본 시스템은 단순 키워드 기반 검토에서 나아가 기준 시방서 유사 문장을 함께 제시함으로써 검토 결과의 근거성을 보완하였다.")
    lines.append("다만 최종 시방서 수정 시에는 프로젝트 특성, 발주처 기준, 관련 법령 및 최신 표준시방서를 반드시 추가 확인해야 한다.")

    return "\n".join(lines)


def create_docx_report(df, title="시방서 Rule·Retrieval·LLM 검토 보고서"):
    if Document is None:
        raise RuntimeError("python-docx가 설치되어 있지 않습니다.")

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    status_count, trade_count, problem_df = make_summary_tables(df)

    doc.add_heading("1. 검토 개요", level=1)
    doc.add_paragraph(
        "본 보고서는 PDF 또는 TXT 형태의 시방서 문장을 대상으로 Rule Agent, "
        "Embedding Retrieval Agent, LLM Agent를 순차적으로 적용하여 생성하였다. "
        "Rule Agent는 공종별 키워드와 필수항목 포함 여부를 기준으로 누락 및 모호 표현을 판정하고, "
        "Embedding Retrieval Agent는 기준 시방서 문장 중 유사 근거를 검색하며, "
        "LLM Agent는 이를 바탕으로 검토 의견과 수정 예시 문장을 생성한다."
    )

    doc.add_heading("2. 판정 결과 요약", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "판정"
    hdr[1].text = "건수"
    for _, row in status_count.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["판정"])
        cells[1].text = str(row["건수"])

    doc.add_heading("3. 공종별 분류 결과", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "분류"
    hdr[1].text = "건수"
    for _, row in trade_count.head(15).iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["분류"])
        cells[1].text = str(row["건수"])

    doc.add_heading("4. 대표 검토 필요 문장", level=1)

    if problem_df.empty:
        doc.add_paragraph("검토 필요 문장이 발견되지 않았다.")
    else:
        for _, row in problem_df.head(10).iterrows():
            doc.add_heading(f"조항 {row.get('조항번호', '-')}", level=2)
            doc.add_paragraph(f"분류: {row.get('분류', '-')}")
            doc.add_paragraph(f"판정: {row.get('판정', '-')}")
            doc.add_paragraph(f"대상 문장: {row.get('문장', '-')}")
            doc.add_paragraph(f"누락 항목: {row.get('누락 항목', '-')}")
            doc.add_paragraph(f"모호 표현: {row.get('모호 표현', '-')}")
            doc.add_paragraph(f"Rule 조언: {row.get('조언', '-')}")
            if "Retrieval 근거" in row:
                doc.add_paragraph(f"Retrieval 근거: {row.get('Retrieval 근거', '-')}")
            if "LLM 검토 의견" in row:
                doc.add_paragraph("LLM 검토 의견:")
                doc.add_paragraph(str(row.get("LLM 검토 의견", "-")))

    doc.add_heading("5. 종합 의견", level=1)
    doc.add_paragraph(
        "본 시스템은 단순 키워드 기반 검토에서 나아가 기준 시방서 유사 문장을 함께 제시함으로써 "
        "검토 결과의 근거성을 보완하였다. 다만 최종 시방서 수정 시에는 프로젝트 특성, 발주처 기준, "
        "관련 법령 및 최신 표준시방서를 반드시 추가 확인해야 한다."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# =========================================================
# 9. Streamlit UI
# =========================================================

def sidebar_rules_editor():
    st.sidebar.subheader("Rule 설정")

    use_default = st.sidebar.checkbox("기본 10개 공종 Rule 사용", value=True)

    rules = DEFAULT_TRADE_RULES.copy() if use_default else []

    with st.sidebar.expander("사용자 공종 Rule 추가"):
        custom_trade = st.text_input("공종명", placeholder="예: 토공사")
        custom_keywords = st.text_input("키워드, 쉼표 구분", placeholder="예: 터파기, 되메우기, 토사")
        custom_required = st.text_area(
            "필수항목 입력",
            placeholder="예:\n시공 범위:범위|구간|위치\n다짐 기준:다짐|층다짐|밀도\n검사 기준:검사|시험|확인"
        )

        if custom_trade and custom_keywords:
            required_dict = {}
            for line in custom_required.splitlines():
                if ":" in line:
                    key, pattern = line.split(":", 1)
                    required_dict[key.strip()] = pattern.strip()
            if not required_dict:
                required_dict = {"기준 명시": r"(기준|방법|검사|확인|수치|절차)"}

            rules.append({
                "trade": custom_trade.strip(),
                "target_keywords": [x.strip() for x in custom_keywords.split(",") if x.strip()],
                "required": required_dict
            })

    vague_text = st.sidebar.text_area(
        "모호 표현 목록",
        value=", ".join(DEFAULT_VAGUE_WORDS),
        height=100
    )
    vague_words = [x.strip() for x in vague_text.split(",") if x.strip()]

    return rules, vague_words


def render_metrics(df):
    if df is None or df.empty:
        return

    problem_count = df["판정"].isin(["누락 가능성", "모호 표현", "누락 및 모호", "검토 필요"]).sum()

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("전체 문장", len(df))
    c2.metric("검토 필요", int(problem_count))
    c3.metric("누락 가능성", int((df["판정"] == "누락 가능성").sum()))
    c4.metric("누락 및 모호", int((df["판정"] == "누락 및 모호").sum()))


def filter_result_ui(df):
    if df is None or df.empty:
        return df

    with st.expander("결과 필터", expanded=False):
        status_options = ["전체"] + sorted(df["판정"].dropna().unique().tolist())
        status = st.selectbox("판정 필터", status_options)

        trade_options = ["전체"] + sorted(df["분류"].dropna().unique().tolist())
        trade = st.selectbox("분류 필터", trade_options)

        keyword = st.text_input("문장 검색")

    out = df.copy()
    if status != "전체":
        out = out[out["판정"] == status]
    if trade != "전체":
        out = out[out["분류"] == trade]
    if keyword:
        out = out[out["문장"].astype(str).str.contains(keyword, case=False, na=False)]

    return out


def main():
    st.set_page_config(
        page_title="시방서 Agentic AI 검토 시스템",
        layout="wide"
    )

    st.title("시방서 Agentic AI 검토 시스템")
    st.caption("PDF → Rule Agent → Embedding Retrieval Agent → LLM Agent → 보고서 생성")

    rules, vague_words = sidebar_rules_editor()

    st.sidebar.subheader("OpenAI API 설정")
    input_api_key = st.sidebar.text_input(
        "OPENAI_API_KEY",
        type="password",
        help="환경변수 또는 Streamlit secrets에 설정했다면 비워도 됩니다."
    )

    generation_model = st.sidebar.text_input("LLM 모델", value="gpt-5.5")
    embedding_model = st.sidebar.text_input("Embedding 모델", value="text-embedding-3-small")

    mode = st.sidebar.radio(
        "기능 선택",
        [
            "전체 자동 실행",
            "1. PDF → TXT 추출",
            "2. Rule 기반 검토",
            "3. Embedding Retrieval + LLM",
            "4. 보고서 생성"
        ]
    )

    if "clean_text" not in st.session_state:
        st.session_state.clean_text = ""
    if "rule_df" not in st.session_state:
        st.session_state.rule_df = None
    if "ref_index" not in st.session_state:
        st.session_state.ref_index = None
    if "final_df" not in st.session_state:
        st.session_state.final_df = None

    # -----------------------------------------------------
    # 전체 자동 실행
    # -----------------------------------------------------
    if mode == "전체 자동 실행":
        st.header("전체 자동 실행")
        st.write("PDF 시방서와 기준 시방서 TXT를 업로드하면 Rule 검토, Embedding 검색, LLM 검토 의견, 보고서 생성까지 한 번에 진행합니다.")

        pdf_file = st.file_uploader("검토 대상 PDF 업로드", type=["pdf"])
        method = st.radio("PDF 추출 방식", ["pdfplumber", "pymupdf"], horizontal=True)

        reference_files = st.file_uploader(
            "기준 시방서 TXT 업로드",
            type=["txt"],
            accept_multiple_files=True
        )

        c1, c2, c3 = st.columns(3)
        top_k = c1.number_input("검색 근거 문장 수", min_value=1, max_value=10, value=3)
        max_retrieval_rows = c2.number_input("Retrieval 대상 최대 문장 수", min_value=1, max_value=200, value=30)
        max_llm_rows = c3.number_input("LLM 생성 대상 최대 문장 수", min_value=1, max_value=50, value=10)

        run_llm = st.checkbox("LLM 검토 의견 생성", value=True)

        if st.button("전체 자동 실행 시작", type="primary"):
            if pdf_file is None:
                st.error("검토 대상 PDF를 업로드해주세요.")
                return

            api_key = get_openai_api_key(input_api_key)
            if not api_key and (reference_files or run_llm):
                st.error("Embedding 또는 LLM 기능을 사용하려면 OPENAI_API_KEY가 필요합니다.")
                return

            with st.spinner("PDF에서 TXT를 추출하는 중입니다..."):
                clean_text = pdf_to_clean_txt(pdf_file, method)
                st.session_state.clean_text = clean_text

            with st.spinner("Rule Agent가 시방서 문장을 검토하는 중입니다..."):
                rule_df = analyze_spec_text(clean_text, rules, vague_words)
                st.session_state.rule_df = rule_df
                st.session_state.final_df = rule_df

            if reference_files:
                client = get_openai_client(api_key)

                with st.spinner("기준 시방서 문장을 로딩하는 중입니다..."):
                    reference_df = load_reference_texts(reference_files)
                    st.write(f"기준 시방서 문장 수: {len(reference_df)}개")

                with st.spinner("기준 시방서 Embedding Index를 생성하는 중입니다..."):
                    ref_index = build_reference_embedding_index(client, reference_df, embedding_model)
                    st.session_state.ref_index = ref_index

                with st.spinner("Embedding Retrieval Agent가 유사 기준 문장을 검색하는 중입니다..."):
                    final_df = add_embedding_retrieval_to_df(
                        client=client,
                        rule_df=rule_df,
                        ref_index=ref_index,
                        embedding_model=embedding_model,
                        top_k=top_k,
                        max_rows=max_retrieval_rows
                    )
                    st.session_state.final_df = final_df

                if run_llm:
                    with st.spinner("LLM Agent가 검토 의견과 수정 예시 문장을 생성하는 중입니다..."):
                        final_df = add_llm_reviews_to_df(
                            client=client,
                            df=st.session_state.final_df,
                            generation_model=generation_model,
                            max_rows=max_llm_rows
                        )
                        st.session_state.final_df = final_df

            st.success("전체 자동 실행 완료")

        df = st.session_state.final_df
        if df is not None:
            render_metrics(df)
            filtered = filter_result_ui(df)
            st.dataframe(filtered, height=500)

            st.download_button(
                "Excel 결과 다운로드",
                data=make_download_bytes_df(df, "xlsx"),
                file_name=now_name("spec_agent_result", "xlsx"),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

            md = generate_markdown_report(df)
            st.download_button(
                "Markdown 보고서 다운로드",
                data=md.encode("utf-8-sig"),
                file_name=now_name("spec_agent_report", "md"),
                mime="text/markdown"
            )

            try:
                docx_buffer = create_docx_report(df)
                st.download_button(
                    "Word 보고서 다운로드",
                    data=docx_buffer,
                    file_name=now_name("spec_agent_report", "docx"),
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            except Exception as e:
                st.warning(f"Word 보고서 생성 불가: {e}")

    # -----------------------------------------------------
    # 1. PDF → TXT 추출
    # -----------------------------------------------------
    elif mode == "1. PDF → TXT 추출":
        st.header("1. PDF → TXT 추출")

        pdf_file = st.file_uploader("PDF 업로드", type=["pdf"])
        method = st.radio("PDF 추출 방식", ["pdfplumber", "pymupdf"], horizontal=True)

        if st.button("TXT 추출 실행"):
            if pdf_file is None:
                st.error("PDF를 업로드해주세요.")
                return

            with st.spinner("PDF에서 텍스트를 추출하는 중입니다..."):
                clean_text = pdf_to_clean_txt(pdf_file, method)
                st.session_state.clean_text = clean_text

            st.success("TXT 추출 완료")

        if st.session_state.clean_text:
            st.text_area("TXT 미리보기", st.session_state.clean_text[:10000], height=400)
            st.download_button(
                "TXT 다운로드",
                data=st.session_state.clean_text.encode("utf-8-sig"),
                file_name=now_name("cleaned_spec", "txt"),
                mime="text/plain"
            )

    # -----------------------------------------------------
    # 2. Rule 기반 검토
    # -----------------------------------------------------
    elif mode == "2. Rule 기반 검토":
        st.header("2. Rule 기반 검토")

        txt_file = st.file_uploader("TXT 업로드", type=["txt"])
        use_session_text = st.checkbox("이전에 추출한 TXT 사용", value=bool(st.session_state.clean_text))

        if st.button("Rule 검토 실행"):
            if use_session_text and st.session_state.clean_text:
                text = st.session_state.clean_text
            elif txt_file is not None:
                text = read_uploaded_text(txt_file)
            else:
                st.error("TXT 파일을 업로드하거나 이전에 추출한 TXT를 사용해주세요.")
                return

            with st.spinner("Rule Agent가 검토하는 중입니다..."):
                df = analyze_spec_text(text, rules, vague_words)
                st.session_state.rule_df = df
                st.session_state.final_df = df

            st.success("Rule 검토 완료")

        df = st.session_state.rule_df
        if df is not None:
            render_metrics(df)
            filtered = filter_result_ui(df)
            st.dataframe(filtered, height=500)
            st.download_button(
                "Rule 결과 Excel 다운로드",
                data=make_download_bytes_df(df, "xlsx"),
                file_name=now_name("rule_check_result", "xlsx"),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

    # -----------------------------------------------------
    # 3. Embedding Retrieval + LLM
    # -----------------------------------------------------
    elif mode == "3. Embedding Retrieval + LLM":
        st.header("3. Embedding Retrieval + LLM")

        uploaded_excel = st.file_uploader("Rule 결과 Excel 업로드", type=["xlsx"])
        use_session_df = st.checkbox("이전 Rule 결과 사용", value=st.session_state.rule_df is not None)

        reference_files = st.file_uploader(
            "기준 시방서 TXT 업로드",
            type=["txt"],
            accept_multiple_files=True
        )

        c1, c2, c3 = st.columns(3)
        top_k = c1.number_input("검색 근거 문장 수", min_value=1, max_value=10, value=3)
        max_retrieval_rows = c2.number_input("Retrieval 대상 최대 문장 수", min_value=1, max_value=200, value=30)
        max_llm_rows = c3.number_input("LLM 생성 대상 최대 문장 수", min_value=1, max_value=50, value=10)

        run_retrieval = st.checkbox("Embedding Retrieval 실행", value=True)
        run_llm = st.checkbox("LLM 검토 의견 생성", value=True)

        if st.button("Embedding + LLM 실행", type="primary"):
            api_key = get_openai_api_key(input_api_key)
            if not api_key:
                st.error("OPENAI_API_KEY가 필요합니다.")
                return

            if use_session_df and st.session_state.rule_df is not None:
                df = st.session_state.rule_df.copy()
            elif uploaded_excel is not None:
                df = pd.read_excel(uploaded_excel)
            else:
                st.error("Rule 결과 Excel을 업로드하거나 이전 Rule 결과를 사용해주세요.")
                return

            client = get_openai_client(api_key)

            if run_retrieval:
                if not reference_files:
                    st.error("Embedding Retrieval을 실행하려면 기준 시방서 TXT가 필요합니다.")
                    return

                with st.spinner("기준 시방서 문장을 로딩하는 중입니다..."):
                    reference_df = load_reference_texts(reference_files)
                    st.write(f"기준 시방서 문장 수: {len(reference_df)}개")

                with st.spinner("기준 시방서 Embedding Index를 생성하는 중입니다..."):
                    ref_index = build_reference_embedding_index(client, reference_df, embedding_model)
                    st.session_state.ref_index = ref_index

                with st.spinner("유사 기준 문장을 검색하는 중입니다..."):
                    df = add_embedding_retrieval_to_df(
                        client=client,
                        rule_df=df,
                        ref_index=ref_index,
                        embedding_model=embedding_model,
                        top_k=top_k,
                        max_rows=max_retrieval_rows
                    )

            if run_llm:
                with st.spinner("LLM 검토 의견을 생성하는 중입니다..."):
                    df = add_llm_reviews_to_df(
                        client=client,
                        df=df,
                        generation_model=generation_model,
                        max_rows=max_llm_rows
                    )

            st.session_state.final_df = df
            st.success("Embedding + LLM 처리 완료")

        df = st.session_state.final_df
        if df is not None:
            render_metrics(df)
            filtered = filter_result_ui(df)
            st.dataframe(filtered, height=500)

            st.download_button(
                "Embedding+LLM 결과 Excel 다운로드",
                data=make_download_bytes_df(df, "xlsx"),
                file_name=now_name("rule_retrieval_llm_result", "xlsx"),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

    # -----------------------------------------------------
    # 4. 보고서 생성
    # -----------------------------------------------------
    elif mode == "4. 보고서 생성":
        st.header("4. 보고서 생성")

        uploaded_excel = st.file_uploader("최종 결과 Excel 업로드", type=["xlsx"])
        use_session_df = st.checkbox("현재 세션 결과 사용", value=st.session_state.final_df is not None)

        df = None
        if use_session_df and st.session_state.final_df is not None:
            df = st.session_state.final_df
        elif uploaded_excel is not None:
            df = pd.read_excel(uploaded_excel)

        if df is not None:
            render_metrics(df)
            st.dataframe(filter_result_ui(df), height=400)

            md = generate_markdown_report(df)
            st.subheader("보고서 미리보기")
            st.markdown(md[:8000])

            st.download_button(
                "Markdown 보고서 다운로드",
                data=md.encode("utf-8-sig"),
                file_name=now_name("spec_agent_report", "md"),
                mime="text/markdown"
            )

            try:
                docx_buffer = create_docx_report(df)
                st.download_button(
                    "Word 보고서 다운로드",
                    data=docx_buffer,
                    file_name=now_name("spec_agent_report", "docx"),
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            except Exception as e:
                st.warning(f"Word 보고서 생성 불가: {e}")
        else:
            st.info("최종 결과 Excel을 업로드하거나 현재 세션 결과를 사용해주세요.")


if __name__ == "__main__":
    main()
